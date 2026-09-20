"""
Generate a 16-slide PowerPoint deck for a topic from .tmp/<slug>/analysis.json.
Design: dark navy (#0A1628) background, white text, cyan (#00D4FF) accents.

  python tools/generate_report.py --topic animation

Output: .tmp/<slug>/report_<slug>_YYYY-MM-DD.pptx

Dependencies: python-pptx, plotly, kaleido
"""

import json
import os
import sys
from datetime import datetime, timezone

import plotly.graph_objects as go
from plotly.subplots import make_subplots
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
import common  # noqa: E402

TMP_DIR = None  # set in main() to .tmp/<slug>/

# Design constants
BG = RGBColor(0x0A, 0x16, 0x28)       # Dark navy
ACCENT = RGBColor(0x00, 0xD4, 0xFF)   # Cyan
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xAA, 0xAA, 0xAA)
DARK_ACCENT = RGBColor(0x00, 0x8A, 0xA8)
GREEN = RGBColor(0x7E, 0xD3, 0x21)
AMBER = RGBColor(0xF5, 0xA6, 0x23)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

PLOTLY_PAPER_BG = "rgba(10,22,40,1)"
PLOTLY_FONT = dict(family="Arial", color="white", size=13)
PLOTLY_CYAN = "#00D4FF"
PLOTLY_BLUE = "#4A90D9"
PLOTLY_MUTED = "#5A6B7F"


# ---------------------------------------------------------------------------
# Slide helpers
# ---------------------------------------------------------------------------

def set_bg(slide, prs):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, text, left, top, width, height,
             font_size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txb


def add_divider(slide, top):
    line = slide.shapes.add_connector(1, Inches(0.5), top, SLIDE_W - Inches(0.5), top)
    line.line.color.rgb = ACCENT
    line.line.width = Pt(1.5)


def add_slide(prs, layout_idx=6):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    set_bg(slide, prs)
    return slide


def embed_chart(slide, img_path, left=Inches(0.4), top=Inches(1.3),
                width=Inches(12.5), height=Inches(5.8)):
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, left, top, width, height)


def save_plotly(fig, filename, width=1300, height=620):
    path = os.path.join(TMP_DIR, filename)
    fig.write_image(path, width=width, height=height, scale=2)
    return path


def apply_plotly_theme(fig):
    fig.update_layout(
        paper_bgcolor=PLOTLY_PAPER_BG,
        plot_bgcolor=PLOTLY_PAPER_BG,
        font=PLOTLY_FONT,
        margin=dict(l=60, r=40, t=40, b=60),
        legend=dict(bgcolor="rgba(0,0,0,0)", font=dict(color="white")),
    )
    fig.update_xaxes(color="white", gridcolor="rgba(255,255,255,0.1)", tickfont=dict(color="white"))
    fig.update_yaxes(color="white", gridcolor="rgba(255,255,255,0.1)", tickfont=dict(color="white"))
    return fig


def slide_header(slide, title, subtitle=None):
    add_text(slide, title, left=Inches(0.5), top=Inches(0.2), width=Inches(12), height=Inches(0.7),
             font_size=28, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, subtitle, left=Inches(0.5), top=Inches(0.85), width=Inches(12), height=Inches(0.35),
                 font_size=13, color=GRAY)
    add_divider(slide, Inches(1.2))


def trunc(text: str, n: int) -> str:
    text = text or ""
    return text if len(text) <= n else text[: n - 1] + "…"


def add_table(slide, top, cols, widths, rows, row_h=Inches(0.42), font_size=11):
    """Text-based table: cyan header row, then one text box per cell."""
    lefts = [Inches(0.3)]
    for w in widths[:-1]:
        lefts.append(lefts[-1] + w)
    for col, left, w in zip(cols, lefts, widths):
        add_text(slide, col, left=left, top=top, width=w, height=row_h,
                 font_size=12, bold=True, color=ACCENT)
    for i, row in enumerate(rows):
        row_top = top + Inches(0.4) + row_h * i
        for val, left, w in zip(row, lefts, widths):
            add_text(slide, str(val), left=left, top=row_top, width=w, height=row_h,
                     font_size=font_size, color=WHITE)


def add_note(slide, text, top=Inches(6.95), color=GRAY):
    add_text(slide, text, left=Inches(0.5), top=top, width=Inches(12.3), height=Inches(0.4),
             font_size=10, color=color)


# ---------------------------------------------------------------------------
# Individual slides
# ---------------------------------------------------------------------------

def slide_title(prs, data, date_str):
    slide = add_slide(prs)
    add_text(slide, data["display_name"], left=Inches(1), top=Inches(1.5), width=Inches(11), height=Inches(1.2),
             font_size=52, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "YouTube Intelligence Report", left=Inches(1), top=Inches(2.6), width=Inches(11),
             height=Inches(0.9), font_size=36, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, date_str, left=Inches(1), top=Inches(3.6), width=Inches(11), height=Inches(0.5),
             font_size=18, color=GRAY, align=PP_ALIGN.CENTER)
    scope = (f"Analysis of {data['total_videos_analyzed']:,} videos from the last {data.get('lookback_days', 30)} days "
             f"of keyword searches plus {data.get('channels_tracked', 0)} tracked channels")
    add_text(slide, scope, left=Inches(1), top=Inches(4.2), width=Inches(11), height=Inches(0.5),
             font_size=15, color=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "Powered by WAT Framework", left=Inches(1), top=Inches(6.8), width=Inches(11),
             height=Inches(0.4), font_size=11, color=GRAY, align=PP_ALIGN.CENTER)


def slide_executive_summary(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Executive Summary", "Key findings from this week's data")
    for i, insight in enumerate(data.get("key_insights", [])[:6]):
        add_text(slide, f"●  {insight}", left=Inches(0.6), top=Inches(1.5) + Inches(i * 0.85),
                 width=Inches(12.1), height=Inches(0.8), font_size=14, color=WHITE)
    notes = [data.get("history_note", "")] + [f"Warning: {w}" for w in data.get("fetch_warnings", [])[:3]]
    add_note(slide, "  |  ".join(n for n in notes if n), top=Inches(6.75))


def slide_top_videos_table(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Top 10 Videos by Views", f"Most watched {data['display_name']} content right now")
    rows = [[i + 1, trunc(v["title"], 55), trunc(v["channel_name"], 28), f"{v['view_count']:,}",
             f"{v['engagement_rate'] * 100:.1f}%"] for i, v in enumerate(data["top_videos_by_views"][:10])]
    add_table(slide, Inches(1.4),
              ["#", "Title", "Channel", "Views", "Eng. Rate"],
              [Inches(0.4), Inches(6.2), Inches(2.5), Inches(1.5), Inches(1.2)],
              rows, row_h=Inches(0.47))


def slide_most_viewed(prs, data):
    top = data["top_videos_by_views"][:15]
    labels = [trunc(v["title"], 35) for v in top][::-1]
    values = [v["view_count"] for v in top][::-1]
    fig = go.Figure(go.Bar(
        x=values, y=labels, orientation="h", marker_color=PLOTLY_CYAN,
        text=[f"{v:,.0f}" for v in values], textposition="outside", textfont=dict(color="white", size=10),
    ))
    apply_plotly_theme(fig)
    fig.update_layout(xaxis_title="View Count", yaxis_title="")
    slide = add_slide(prs)
    slide_header(slide, "Most Viewed", "Top 15 videos by total view count")
    embed_chart(slide, save_plotly(fig, "chart_most_viewed.png"))


def slide_engagement_champions(prs, data):
    top = data["top_videos_by_engagement"][:10]
    labels = [trunc(v["title"], 35) for v in top][::-1]
    values = [round(v["engagement_rate"] * 100, 2) for v in top][::-1]
    fig = go.Figure(go.Bar(
        x=values, y=labels, orientation="h", marker_color="#FF6B6B",
        text=[f"{v:.1f}%" for v in values], textposition="outside", textfont=dict(color="white", size=11),
    ))
    apply_plotly_theme(fig)
    fig.update_layout(xaxis_title="Engagement Rate (%)", yaxis_title="")
    slide = add_slide(prs)
    slide_header(slide, "Engagement Champions", "Highest (likes + comments) / views ratio, videos over 1,000 views")
    embed_chart(slide, save_plotly(fig, "chart_engagement.png"))


def slide_trending_phrases(prs, data):
    phrases = data.get("trending_phrases", [])[:15]
    if phrases:
        labels = [p["phrase"] for p in phrases][::-1]
        values = [p["video_count"] for p in phrases][::-1]
        subtitle = "Two-word phrases in titles, by number of videos using them"
    else:
        kws = data.get("trending_keywords", [])[:15]
        labels = [k["keyword"] for k in kws][::-1]
        values = [k["count"] for k in kws][::-1]
        subtitle = "Most frequent title words (not enough data for two-word phrases)"
    colors = [PLOTLY_CYAN if i >= len(values) - 5 else PLOTLY_BLUE for i in range(len(values))]
    fig = go.Figure(go.Bar(
        x=values, y=labels, orientation="h", marker_color=colors,
        text=values, textposition="outside", textfont=dict(color="white", size=11),
    ))
    apply_plotly_theme(fig)
    fig.update_layout(xaxis_title="Videos", yaxis_title="")
    slide = add_slide(prs)
    slide_header(slide, "Trending Phrases", subtitle)
    embed_chart(slide, save_plotly(fig, "chart_phrases.png"))


def slide_emerging(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Emerging Topics", "Phrases that are new or growing versus the previous snapshot")
    emerging = data.get("emerging_topics", [])
    if not emerging:
        add_text(slide, data.get("history_note", ""), left=Inches(0.6), top=Inches(1.6), width=Inches(12),
                 height=Inches(0.8), font_size=16, color=WHITE)
        add_text(slide,
                 "Each weekly run stores its phrase counts. From the second run on, this slide lists phrases "
                 "that are new or whose share of videos grew by at least 50%.",
                 left=Inches(0.6), top=Inches(2.6), width=Inches(12), height=Inches(1.2), font_size=14, color=GRAY)
        return
    rows = []
    for e in emerging[:8]:
        change = "NEW" if e["status"] == "new" else f"{e['change']:.1f}x"
        top_v = (e.get("top_videos") or [{}])[0]
        related = f"{trunc(top_v.get('title', ''), 52)} ({top_v.get('channel_name', '')[:18]})" if top_v else ""
        rows.append([trunc(e["phrase"], 26), change, f"{e['video_count']} (was {e['previous_count']})", related])
    add_table(slide, Inches(1.4), ["Phrase", "Change", "Videos now", "Top related video"],
              [Inches(2.8), Inches(1.0), Inches(1.9), Inches(7.2)], rows, row_h=Inches(0.5))
    cooling = data.get("cooling_topics", [])
    if cooling:
        add_note(slide, "Cooling: " + ", ".join(f"{c['phrase']} ({c['previous_count']} -> {c['video_count']})"
                                                for c in cooling[:5]), top=Inches(6.7))
    add_note(slide, data.get("history_note", ""), top=Inches(7.0))


def slide_clusters(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Topic Clusters & Related Videos",
                 "Videos grouped by similar titles and tags, with the best performers in each group")
    clusters = data.get("clusters", [])
    if not clusters:
        add_text(slide, f"Clusters not built: {data.get('clusters_skipped_reason') or 'no data'}.",
                 left=Inches(0.6), top=Inches(1.6), width=Inches(12), height=Inches(0.8), font_size=16)
        return
    rows = []
    for c in clusters[:8]:
        top_v = c["top_videos"][0] if c["top_videos"] else {}
        rows.append([trunc(c["label"], 34), c["video_count"],
                     f"{c['median_outlier']:.1f}x" if c["median_outlier"] is not None else "n/a",
                     f"{trunc(top_v.get('title', ''), 58)} ({trunc(top_v.get('channel_name', ''), 20)})"])
    add_table(slide, Inches(1.4), ["Topic cluster", "Videos", "vs. norm", "Best-performing video"],
              [Inches(3.5), Inches(0.9), Inches(1.1), Inches(7.4)], rows, row_h=Inches(0.55))
    add_note(slide, "vs. norm: median views relative to each channel's typical video (1.0x = typical).")


def slide_skills(prs, data):
    skills = [s for s in data.get("skills", []) if s["video_count"] > 0][:12]
    slide = add_slide(prs)
    slide_header(slide, "Skills & Tools",
                 "Bar = videos mentioning it; label = median performance vs channel norm (1.0x = typical)")
    if not skills:
        add_text(slide, "No skill terms matched any video in this sample.", left=Inches(0.6), top=Inches(1.6),
                 width=Inches(12), height=Inches(0.8), font_size=16)
        return
    skills = skills[::-1]
    labels = [f"{s['skill']}  ({s['median_outlier']:.1f}x)" if s["median_outlier"] is not None
              else f"{s['skill']}  (n/a)" for s in skills]
    values = [s["video_count"] for s in skills]
    colors = [PLOTLY_MUTED if s["median_outlier"] is None
              else PLOTLY_CYAN if s["median_outlier"] >= 1.2 else PLOTLY_BLUE for s in skills]
    fig = go.Figure(go.Bar(x=values, y=labels, orientation="h", marker_color=colors,
                           text=values, textposition="outside", textfont=dict(color="white", size=11)))
    apply_plotly_theme(fig)
    fig.update_layout(xaxis_title="Videos mentioning the skill", yaxis_title="")
    embed_chart(slide, save_plotly(fig, "chart_skills.png"))
    add_note(slide, "Cyan = overperforming (1.2x or better). Grey = too few videos from channels with enough history "
                    "to measure performance.")


def slide_channel_comparison(prs, data):
    channels = [c for c in data.get("channel_stats", []) if c["video_count"] >= 3][:10]
    names = [trunc(c["channel_name"], 22) for c in channels]
    fig = go.Figure()
    fig.add_trace(go.Bar(name="Avg Views/Video", x=names, y=[c["avg_views"] for c in channels],
                         marker_color=PLOTLY_CYAN, yaxis="y"))
    fig.add_trace(go.Bar(name="Videos Analyzed", x=names, y=[c["video_count"] for c in channels],
                         marker_color="#FF6B6B", yaxis="y2"))
    apply_plotly_theme(fig)
    fig.update_layout(
        barmode="group",
        yaxis=dict(title="Avg Views/Video", color="white", gridcolor="rgba(255,255,255,0.1)"),
        yaxis2=dict(title="Video Count", overlaying="y", side="right", color="white"),
        legend=dict(orientation="h", y=1.1), xaxis_tickangle=-30,
    )
    slide = add_slide(prs)
    slide_header(slide, "Channel Comparison", "Top channels by total views (3+ videos in the sample)")
    embed_chart(slide, save_plotly(fig, "chart_channels.png"))


def slide_channels_to_watch(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Channels to Watch",
                 "Channels not in your list that keep showing up in topic searches")
    cands = data.get("channel_candidates", [])[:10]
    if not cands:
        add_text(slide, "No new channel candidates this run.", left=Inches(0.6), top=Inches(1.6),
                 width=Inches(12), height=Inches(0.8), font_size=16)
        return
    rows = []
    for c in cands:
        subs = f"{c['subscribers']:,}" if c.get("subscribers") else "n/a"
        vps = f"{c['views_per_subscriber']:.1f}" if c.get("views_per_subscriber") is not None else "n/a"
        rows.append([trunc(c["channel_name"], 36), subs, c["video_count"], f"{c['median_views']:,}", vps])
    add_table(slide, Inches(1.4), ["Channel", "Subscribers", "Videos seen", "Median views", "Views / subscriber"],
              [Inches(4.6), Inches(2.0), Inches(1.6), Inches(2.1), Inches(2.3)], rows, row_h=Inches(0.48))
    add_note(slide, "High views per subscriber means a video reaches well beyond the channel's own audience. "
                    "To track one, add its handle to the topic profile.")


def slide_content_and_format(prs, data):
    dist = data.get("content_type_distribution", {})
    colors = ["#00D4FF", "#FF6B6B", "#4A90D9", "#F5A623", "#7ED321", "#9B59B6", "#1ABC9C", "#95A5A6"]
    fig = go.Figure(go.Pie(
        labels=list(dist.keys()), values=list(dist.values()),
        marker=dict(colors=colors[:len(dist)], line=dict(color="#0A1628", width=2)),
        textinfo="label+percent", textfont=dict(size=13, color="white"), hole=0.35,
    ))
    apply_plotly_theme(fig)
    fig.update_layout(showlegend=False)
    slide = add_slide(prs)
    slide_header(slide, "Content Type & Format",
                 f"What kinds of {data['display_name']} videos dominate, and long-form vs Shorts")
    embed_chart(slide, save_plotly(fig, "chart_content_types.png", width=900, height=620),
                left=Inches(0.3), top=Inches(1.4), width=Inches(7.4), height=Inches(5.5))
    fmt = data.get("format_split", {})
    y = Inches(1.7)
    for label, key in (("Long-form", "long"), ("Shorts", "short")):
        f = fmt.get(key, {})
        add_text(slide, label, left=Inches(8.1), top=y, width=Inches(4.8), height=Inches(0.5),
                 font_size=20, bold=True, color=ACCENT)
        add_text(slide, f"{f.get('video_count', 0)} videos", left=Inches(8.1), top=y + Inches(0.5),
                 width=Inches(4.8), height=Inches(0.4), font_size=14)
        add_text(slide, f"Median views: {f.get('median_views', 0):,}", left=Inches(8.1), top=y + Inches(0.9),
                 width=Inches(4.8), height=Inches(0.4), font_size=14)
        add_text(slide, f"Avg engagement: {f.get('avg_engagement_rate', 0) * 100:.1f}%", left=Inches(8.1),
                 top=y + Inches(1.3), width=Inches(4.8), height=Inches(0.4), font_size=14)
        y += Inches(2.4)


def slide_best_day_and_length(prs, data):
    by_day = data.get("performance_by_weekday", [])
    by_len = data.get("performance_by_duration", [])

    def series(rows):
        valid = [r["median_outlier"] for r in rows if r["sufficient"]]
        best = max(valid) if valid else None
        vals = [r["median_outlier"] if r["sufficient"] else 0 for r in rows]
        cols = [PLOTLY_CYAN if (r["sufficient"] and r["median_outlier"] == best)
                else PLOTLY_BLUE if r["sufficient"] else PLOTLY_MUTED for r in rows]
        text = [f"{r['median_outlier']:.1f}x (n={r['n_relative']})" if r["sufficient"]
                else f"n={r.get('n_relative', r['video_count'])}" for r in rows]
        return vals, cols, text

    fig = make_subplots(rows=1, cols=2, subplot_titles=("By publish day (UTC, long-form)", "By video length"))
    for col, rows in ((1, by_day), (2, by_len)):
        vals, cols, text = series(rows)
        fig.add_trace(go.Bar(x=[r["group"] for r in rows], y=vals, marker_color=cols, text=text,
                             textposition="outside", textfont=dict(color="white", size=10)), row=1, col=col)
        fig.add_hline(y=1.0, line_dash="dot", line_color="rgba(255,255,255,0.5)", row=1, col=col)
    apply_plotly_theme(fig)
    fig.update_layout(showlegend=False, yaxis_title="Median views vs channel norm")
    fig.update_annotations(font=dict(color="white", size=14))
    slide = add_slide(prs)
    slide_header(slide, "Best Day & Length to Publish",
                 "Median performance vs each channel's typical video (dotted line = typical). Grey = under 5 videos")
    embed_chart(slide, save_plotly(fig, "chart_day_length.png"))


def slide_breakout_and_rising(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Breakout & Rising Videos",
                 "Overperformers vs their channel's norm, and the fastest movers of the last 7 days")
    breakout = [[trunc(v["title"], 52), trunc(v["channel_name"], 24), f"{v['outlier_score']:.1f}x",
                 f"{v['view_count']:,}"] for v in data.get("breakout_videos", [])[:5]]
    add_table(slide, Inches(1.35), ["Breakout (min 2 days old)", "Channel", "vs. norm", "Views"],
              [Inches(6.5), Inches(3.0), Inches(1.4), Inches(1.6)], breakout, row_h=Inches(0.4))
    rising = [[trunc(v["title"], 52), trunc(v["channel_name"], 24), f"{v['views_per_day']:,.0f}",
               f"{v['days_since_published']:.1f}d"] for v in data.get("rising_fast", [])[:5]]
    if rising:
        add_table(slide, Inches(4.05), ["Rising fast (last 7 days)", "Channel", "Views/day", "Age"],
                  [Inches(6.5), Inches(3.0), Inches(1.4), Inches(1.6)], rising, row_h=Inches(0.4))
    else:
        add_text(slide, "No videos from the last 7 days in this sample.", left=Inches(0.3), top=Inches(4.1),
                 width=Inches(12), height=Inches(0.5), font_size=13, color=GRAY)


def slide_views_vs_engagement(prs, data):
    vids = [v for v in data["videos"] if v["view_count"] > 0][:400]
    fig = go.Figure(go.Scatter(
        x=[v["view_count"] for v in vids], y=[round(v["engagement_rate"] * 100, 2) for v in vids],
        mode="markers",
        marker=dict(color=PLOTLY_CYAN, size=8, opacity=0.6, line=dict(color="white", width=0.5)),
        text=[f"{trunc(v['title'], 40)}<br>{v['channel_name']}" for v in vids],
        hovertemplate="%{text}<br>Views: %{x:,}<br>Eng Rate: %{y:.1f}%<extra></extra>",
    ))
    apply_plotly_theme(fig)
    fig.update_layout(xaxis_title="Total Views", yaxis_title="Engagement Rate (%)", xaxis_type="log")
    slide = add_slide(prs)
    slide_header(slide, "Views vs. Engagement", "Find the sweet spot: high views AND high engagement")
    embed_chart(slide, save_plotly(fig, "chart_scatter.png"))


def build_recommendations(data: dict) -> list[str]:
    """Data-driven action items; each is skipped when the sample is too small to support it."""
    recs = []
    phrases = [p["phrase"] for p in data.get("trending_phrases", [])[:3]] or \
              [k["keyword"] for k in data.get("trending_keywords", [])[:3]]
    emerging = [e["phrase"] for e in data.get("emerging_topics", []) if e["status"] == "new"][:2]
    if phrases:
        line = f"Work these phrases into titles: {', '.join(phrases)}"
        if emerging:
            line += f". Newly appearing: {', '.join(emerging)}"
        recs.append(line)

    lift = 1.2  # only recommend groups that beat channel norms by 20%+; smaller gaps are noise
    types = [r for r in data.get("performance_by_content_type", [])
             if r["sufficient"] and r["group"] != "other" and (r["median_outlier"] or 0) >= lift]
    if types:
        best = max(types, key=lambda r: r["median_outlier"])
        recs.append(f"Lead with {best['group']} content: it runs at {best['median_outlier']:.1f}x channel norms "
                    f"(n={best['n_relative']})")

    lengths = [r for r in data.get("performance_by_duration", []) if r["sufficient"] and r["median_outlier"] >= lift]
    if lengths:
        best = max(lengths, key=lambda r: r["median_outlier"])
        recs.append(f"Target {best['group']} videos: best median performance at {best['median_outlier']:.1f}x "
                    f"channel norms (n={best['n_relative']})")

    days = [r for r in data.get("performance_by_weekday", []) if r["sufficient"] and r["median_outlier"] >= lift]
    if days:
        best = max(days, key=lambda r: r["median_outlier"])
        recs.append(f"Publish on {best['group']}: videos posted that day perform {best['median_outlier']:.1f}x "
                    f"channel norms (UTC, n={best['n_relative']})")

    skills = [s for s in data.get("skills", [])
              if s["video_count"] >= 5 and (s["median_outlier"] or 0) >= lift]
    if skills:
        best = max(skills, key=lambda s: s["median_outlier"])
        recs.append(f"Skill with the strongest pull: {best['skill']} ({best['median_outlier']:.1f}x norm across "
                    f"{best['video_count']} videos)")

    if data.get("breakout_videos"):
        v = data["breakout_videos"][0]
        recs.append(f"Study the breakout \"{trunc(v['title'], 55)}\" ({v['outlier_score']:.1f}x): "
                    "match its hook and title structure")

    fmt = data.get("format_split", {})
    if fmt.get("short", {}).get("video_count", 0) >= 5 and fmt.get("long", {}).get("video_count", 0) >= 5:
        ratio = fmt["long"]["median_views"] / max(fmt["short"]["median_views"], 1)
        recs.append(f"Long-form draws {ratio:.1f}x the median views of Shorts here; use Shorts for reach, "
                    "long-form for depth")
    return recs


def slide_recommendations(prs, data):
    slide = add_slide(prs)
    slide_header(slide, "Content Recommendations", "Data-driven action items, each backed by this run's numbers")
    recs = build_recommendations(data)
    if not recs:
        recs = ["Not enough data in this sample to support recommendations. Widen the topic profile and re-run."]
    for i, rec in enumerate(recs[:7]):
        add_text(slide, f"→  {rec}", left=Inches(0.6), top=Inches(1.5) + Inches(i * 0.75),
                 width=Inches(12.1), height=Inches(0.7), font_size=14, color=WHITE)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    global TMP_DIR
    slug = common.parse_topic_arg()
    apath = common.analysis_path(slug)
    if not os.path.exists(apath):
        raise FileNotFoundError(f"Not found: {apath}. Run analyze_trends.py --topic {slug} first.")
    with open(apath, encoding="utf-8") as f:
        data = json.load(f)
    TMP_DIR = common.tmp_dir(slug)

    now = datetime.now(timezone.utc)
    output_path = os.path.join(TMP_DIR, f"report_{slug}_{now.strftime('%Y-%m-%d')}.pptx")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    steps = [
        ("Title", lambda: slide_title(prs, data, now.strftime("%B %d, %Y"))),
        ("Executive Summary", lambda: slide_executive_summary(prs, data)),
        ("Top Videos Table", lambda: slide_top_videos_table(prs, data)),
        ("Most Viewed", lambda: slide_most_viewed(prs, data)),
        ("Engagement Champions", lambda: slide_engagement_champions(prs, data)),
        ("Trending Phrases", lambda: slide_trending_phrases(prs, data)),
        ("Emerging Topics", lambda: slide_emerging(prs, data)),
        ("Topic Clusters", lambda: slide_clusters(prs, data)),
        ("Skills & Tools", lambda: slide_skills(prs, data)),
        ("Channel Comparison", lambda: slide_channel_comparison(prs, data)),
        ("Channels to Watch", lambda: slide_channels_to_watch(prs, data)),
        ("Content Type & Format", lambda: slide_content_and_format(prs, data)),
        ("Best Day & Length", lambda: slide_best_day_and_length(prs, data)),
        ("Breakout & Rising", lambda: slide_breakout_and_rising(prs, data)),
        ("Views vs Engagement", lambda: slide_views_vs_engagement(prs, data)),
        ("Recommendations", lambda: slide_recommendations(prs, data)),
    ]
    print("Building slides...")
    for i, (name, build) in enumerate(steps, 1):
        print(f"  Slide {i}: {name}")
        build()

    prs.save(output_path)
    print(f"\nDeck saved: {output_path}")
    return output_path


if __name__ == "__main__":
    main()
